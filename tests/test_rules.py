import tempfile
import unittest
from pathlib import Path

from docx import Document

import app


class NamingRulesTests(unittest.TestCase):
    def test_filename_uses_plan_and_month_from_word(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "input.docx"
            doc = Document()
            doc.add_paragraph("Agosto de 2026")
            doc.add_paragraph("Raio-X de Investimentos")
            doc.add_paragraph("Plano CD CPFL")
            doc.save(src)
            plan, month = app.infer_plan_and_month(src, app.DEFAULTS.copy())
            self.assertEqual(plan, "CD CPFL")
            self.assertEqual(month, "agosto de 2026")
            self.assertEqual(
                app.output_filename(src, app.DEFAULTS.copy()),
                "CD CPFL - Raio X de Agosto.pdf",
            )

    def test_override_values_win(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "sample.docx"
            doc = Document()
            doc.add_paragraph("Plano Something")
            doc.save(src)
            cfg = app.DEFAULTS.copy()
            cfg["plan_name_override"] = "Plano escolhido"
            cfg["competency_override"] = "setembro de 2026"
            self.assertEqual(
                app.output_filename(src, cfg),
                "Plano escolhido - Raio X de Setembro.pdf",
            )

    def test_numeric_month_in_source_filename_is_detected(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "FORD_8_2026_Raio-X_OK.docx"
            doc = Document()
            doc.add_paragraph("Plano FORD")
            doc.save(src)
            plan, month = app.infer_plan_and_month(src, app.DEFAULTS.copy())
            self.assertEqual(plan, "FORD")
            self.assertEqual(month, "agosto de 2026")
            self.assertEqual(app.output_filename(src, app.DEFAULTS.copy()), "FORD - Raio X de Agosto.pdf")

    def test_plan_and_month_in_reference_cover_are_detected(self):
        with tempfile.TemporaryDirectory() as tmp:
            src = Path(tmp) / "AbrappPrev_8_2026_OK.docx"
            doc = Document()
            doc.add_paragraph("Agosto de 2026")
            doc.add_paragraph("Raio - X de Investimentos")
            doc.add_paragraph("Plano Abrapp Prev")
            doc.save(src)
            self.assertEqual(app.infer_plan_and_month(src, app.DEFAULTS.copy()), ("Abrapp Prev", "agosto de 2026"))

    def test_invalid_filename_characters_are_sanitized(self):
        self.assertEqual(app.safe_filename("Plano PSAP/Auren"), "Plano PSAP-Auren")


if __name__ == "__main__":
    unittest.main()
